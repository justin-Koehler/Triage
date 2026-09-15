"""Aufwand aus einer geteilten Google-Sheets-Tabelle lesen."""

from __future__ import annotations

import csv
import io
import re
import uuid
from pathlib import Path
from urllib.parse import parse_qs, urlparse

import httpx
from sqlalchemy.orm import Session

from app.config import ROOT
from app.domain.calc import _cit_total, _scs_total, load_rates, parse_number
from app.models import AppSetting

TEMPLATE_PATH = ROOT / "config" / "effort_sheet_template.csv"
DUMMY_TEMPLATE_PATH = ROOT / "config" / "effort_sheet_dummy.csv"
DUMMY_TEMPLATE_URL = (
    "https://docs.google.com/spreadsheets/d/dummyCRITRaufwandTemplate/edit"
)
DUMMY_OPEN_URL = "/effort-sheet"
_SHEET_ID = re.compile(r"/spreadsheets/d/([a-zA-Z0-9_-]+)")
_ALLOWED_HOSTS = {"docs.google.com"}
_NOT_SHARED = "Sheet nicht freigegeben. Unter Freigabe: Jeder mit dem Link (Lesen)."
_SHARE_PREFIX = "effort.share."
_DUMMY_MARK = "dummycritraufwand"


class EffortSheetError(ValueError):
    """URL ungültig oder Sheet nicht lesbar."""


def extract_sheet_id(url: str) -> str | None:
    match = _SHEET_ID.search(url or "")
    return match.group(1) if match else None


def is_dummy_template_url(url: str) -> bool:
    return _DUMMY_MARK in (url or "").lower()


def dummy_open_url(template_url: str) -> str:
    if is_dummy_template_url(template_url) or not extract_sheet_id(template_url):
        return DUMMY_OPEN_URL
    return (template_url or "").strip()


def copy_url(template_url: str) -> str:
    if is_dummy_template_url(template_url):
        return DUMMY_OPEN_URL
    sheet_id = extract_sheet_id(template_url)
    if not sheet_id:
        return ""
    return f"https://docs.google.com/spreadsheets/d/{sheet_id}/copy"


def dummy_template_bytes() -> bytes:
    path = Path(DUMMY_TEMPLATE_PATH)
    if path.exists():
        return path.read_bytes()
    return b"Aufwand FB,Aufwand IT,Summe\n,,\n"


def parse_dummy_template(url: str = "") -> dict[str, str]:
    parsed = parse_effort_csv(dummy_template_bytes().decode("utf-8"))
    parsed["effort_sheet_url"] = (url or DUMMY_TEMPLATE_URL).strip()
    return parsed


def export_csv_url(url: str) -> str:
    parsed = urlparse((url or "").strip())
    host = (parsed.hostname or "").lower()
    if host not in _ALLOWED_HOSTS:
        raise EffortSheetError("Nur Google-Sheets-Links (docs.google.com) sind erlaubt.")
    sheet_id = extract_sheet_id(url)
    if not sheet_id:
        raise EffortSheetError("Keine Spreadsheet-ID in der URL.")
    gid = "0"
    query = parse_qs(parsed.query)
    if query.get("gid"):
        gid = query["gid"][0]
    elif parsed.fragment.startswith("gid="):
        gid = parsed.fragment.split("=", 1)[-1]
    return f"https://docs.google.com/spreadsheets/d/{sheet_id}/export?format=csv&gid={gid}"


def _norm(header: str) -> str:
    return (
        str(header or "")
        .strip()
        .lower()
        .replace("ä", "a")
        .replace("ö", "o")
        .replace("ü", "u")
        .replace("ß", "ss")
    )


def _phase(raw: str) -> str:
    text = _norm(raw)
    if "betrieb" in text or text in {"operate", "ops"}:
        return "operate"
    return "concept"


def _party(raw: str) -> str:
    text = _norm(raw)
    if any(token in text for token in ("cit", "it", "sit")):
        return "cit"
    return "scs"


def _zeit_key(name: str) -> str:
    return " ".join(_norm(name).replace("_", " ").split())


def _parse_fb_it_csv(rows: list[dict], index: dict[str, str]) -> dict[str, str] | None:
    fb_col = next(
        (
            index[k]
            for k in index
            if _zeit_key(k) in {"aufwand fb", "aufwandfb", "zeit 1", "zeit1"}
        ),
        None,
    )
    it_col = next(
        (
            index[k]
            for k in index
            if _zeit_key(k) in {"aufwand it", "aufwandit", "zeit 2", "zeit2"}
        ),
        None,
    )
    if not fb_col and not it_col:
        return None
    fb = 0.0
    it = 0.0
    for row in rows:
        if not any(str(value or "").strip() for value in row.values()):
            continue
        if fb_col:
            fb += parse_number(row.get(fb_col))
        if it_col:
            it += parse_number(row.get(it_col))
    total = fb + it
    return {
        "effort_fb": _format_pt(fb),
        "effort_it": _format_pt(it),
        "concept_scs_pt": _qty(fb),
        "concept_cit_pt": _qty(it),
        "operate_scs_pt": "",
        "operate_cit_pt": "",
        "summe": _qty(total),
        "costs": "",
    }


def share_id_from_url(url: str) -> str | None:
    match = re.search(r"/aufwand/([0-9a-fA-F-]{36})", url or "")
    if not match:
        return None
    try:
        return str(uuid.UUID(match.group(1)))
    except ValueError:
        return None


def _parse_kalkulation_csv(rows: list[dict], index: dict[str, str]) -> dict[str, str] | None:
    art_col = next((index[k] for k in index if k in {"art", "typ"}), None)
    if not art_col:
        return None
    phase_col = next((index[k] for k in index if k in {"phase", "abschnitt"}), None)
    party_col = next((index[k] for k in index if k in {"rolle", "bereich", "partei", "team"}), None)
    qty_col = next(
        (index[k] for k in index if k in {"menge", "pt", "personentage", "betrag"}),
        None,
    )
    desc_col = next(
        (index[k] for k in index if k in {"beschreibung", "aufschluesselung", "text"}),
        None,
    )
    pts = {"concept_scs": 0.0, "concept_cit": 0.0, "operate_scs": 0.0, "operate_cit": 0.0}
    mats = {"concept_scs": 0.0, "concept_cit": 0.0, "operate_scs": 0.0, "operate_cit": 0.0}
    details: dict[str, list[str]] = {key: [] for key in pts}
    for row in rows:
        if not any(str(value or "").strip() for value in row.values()):
            continue
        art = _norm(row.get(art_col))
        if art in {"summe", "brutto", "satz"}:
            continue
        phase = _phase(row.get(phase_col) if phase_col else "")
        party = _party(row.get(party_col) if party_col else "")
        key = f"{phase}_{party}"
        qty = parse_number(row.get(qty_col) if qty_col else "")
        desc = str(row.get(desc_col) if desc_col else "").strip()
        if art in {"sach", "lizenz", "kosten", "material"}:
            mats[key] += qty
            if desc:
                details[key].append(f"{desc} ({_qty(qty)} €)" if qty else desc)
        else:
            pts[key] += qty
            if desc:
                details[key].append(f"{_qty(qty)} PT {desc}".strip() if qty else desc)
    scs = pts["concept_scs"] + pts["operate_scs"]
    cit = pts["concept_cit"] + pts["operate_cit"]
    costs = sum(mats.values())
    return {
        "effort_fb": _format_pt(scs),
        "effort_it": _format_pt(cit),
        "concept_scs_pt": _qty(pts["concept_scs"]),
        "concept_cit_pt": _qty(pts["concept_cit"]),
        "operate_scs_pt": _qty(pts["operate_scs"]),
        "operate_cit_pt": _qty(pts["operate_cit"]),
        "concept_scs_pt_detail": "; ".join(details["concept_scs"]),
        "concept_cit_pt_detail": "; ".join(details["concept_cit"]),
        "operate_scs_pt_detail": "; ".join(details["operate_scs"]),
        "operate_cit_pt_detail": "; ".join(details["operate_cit"]),
        "concept_scs_material": _qty(mats["concept_scs"]),
        "concept_cit_material": _qty(mats["concept_cit"]),
        "operate_scs_material": _qty(mats["operate_scs"]),
        "operate_cit_material": _qty(mats["operate_cit"]),
        "costs": _qty(costs) if costs else "",
        "summe": _qty(scs + cit),
    }


def _csv_rows(text: str) -> list[list[str]]:
    blob = (text or "").replace("\ufeff", "")
    sample = blob[:400]
    delim = ";" if sample.count(";") >= sample.count(",") else ","
    return list(csv.reader(io.StringIO(blob), delimiter=delim))


def _looks_like_vorlage(rows: list[list[str]]) -> bool:
    blob = " ".join(" ".join(row) for row in rows[:40]).lower()
    return "kalkulation" in blob and "personentage" in blob


def _parse_vorlage_csv(rows: list[list[str]]) -> dict[str, str]:
    phase = "concept"
    party = "scs"
    mode = ""
    items: list[tuple[str, str, str, float, str]] = []
    skip_heads = (
        "personentage",
        "sach-/lizenzkosten",
        "sach-/lizenz",
        "aufschluesselung",
        "kosten-kalkulation",
        "gesamt-pt",
        "gesamt-personentage",
        "tarifsatz",
        "tagessatz",
        "freigegebene",
        "summe pt",
        "gemeinkosten",
        "gewinnzuschlag",
        "sachkosten summe",
        "sonstiges",
        "saetze",
        "teilsummen",
        "berechnungen",
    )
    scs_brutto = 0.0
    cit_brutto = 0.0
    for raw in rows:
        cells = [str(c or "").strip() for c in raw]
        if not any(cells):
            continue
        joined = _norm(" ".join(cells))
        first = _norm(cells[0])
        if "betriebs-phase" in joined or first == "betrieb":
            phase = "operate"
        elif "konzeption" in joined:
            phase = "concept"
        if first.startswith("kalkulation it"):
            party = "cit"
            continue
        if first.startswith("kalkulation scs"):
            party = "scs"
            continue
        if first.startswith("personentage"):
            mode = "pt"
            continue
        if first.startswith("sach") and "lizenz" in first + joined:
            mode = "sach"
            continue
        if first.startswith("kosten-kalkulation") or first in {
            "tarifsatz (scs)",
            "tagessatz pt",
        }:
            mode = "calc"
            continue
        if first.startswith("summe brutto"):
            amount = parse_number(cells[-1] if len(cells) > 1 else "")
            if party == "cit":
                cit_brutto += amount
            else:
                scs_brutto += amount
            mode = "calc"
            continue
        if mode not in {"pt", "sach"}:
            continue
        if first.startswith("+") or any(first.startswith(head) for head in skip_heads):
            continue
        qty = parse_number(cells[0])
        desc = cells[1] if len(cells) > 1 else ""
        if not qty and not desc:
            continue
        if _norm(desc).startswith("bitte trage") or _norm(desc).startswith("gesamt-personentage"):
            continue
        if _norm(desc).startswith("aufschluesselung"):
            continue
        items.append((phase, party, mode, qty, desc))
    pts = {"concept_scs": 0.0, "concept_cit": 0.0, "operate_scs": 0.0, "operate_cit": 0.0}
    mats = {key: 0.0 for key in pts}
    details: dict[str, list[str]] = {key: [] for key in pts}
    for phase, party, art, qty, desc in items:
        key = f"{phase}_{party}"
        if art == "sach":
            mats[key] += qty
            if desc:
                details[key].append(f"{desc} ({_qty(qty)} €)" if qty else desc)
        else:
            pts[key] += qty
            if desc:
                details[key].append(f"{_qty(qty)} PT {desc}".strip() if qty else desc)
    scs = pts["concept_scs"] + pts["operate_scs"]
    cit = pts["concept_cit"] + pts["operate_cit"]
    costs = sum(mats.values())
    return {
        "effort_fb": _format_pt(scs),
        "effort_it": _format_pt(cit),
        "concept_scs_pt": _qty(pts["concept_scs"]),
        "concept_cit_pt": _qty(pts["concept_cit"]),
        "operate_scs_pt": _qty(pts["operate_scs"]),
        "operate_cit_pt": _qty(pts["operate_cit"]),
        "concept_scs_pt_detail": "; ".join(details["concept_scs"]),
        "concept_cit_pt_detail": "; ".join(details["concept_cit"]),
        "operate_scs_pt_detail": "; ".join(details["operate_scs"]),
        "operate_cit_pt_detail": "; ".join(details["operate_cit"]),
        "concept_scs_material": _qty(mats["concept_scs"]),
        "concept_cit_material": _qty(mats["concept_cit"]),
        "operate_scs_material": _qty(mats["operate_scs"]),
        "operate_cit_material": _qty(mats["operate_cit"]),
        "costs": _qty(scs_brutto) if scs_brutto else _qty(costs),
        "it_costs": _qty(cit_brutto) if cit_brutto else "",
        "summe": _qty(scs + cit),
    }


def parse_effort_csv(text: str) -> dict[str, str]:
    raw_rows = _csv_rows(text)
    if _looks_like_vorlage(raw_rows):
        return _with_it_costs(_parse_vorlage_csv(raw_rows))
    reader = csv.DictReader(io.StringIO((text or "").replace("\ufeff", "")))
    if not reader.fieldnames:
        raise EffortSheetError("Tabelle ohne Kopfzeile.")
    index = {_norm(name): name for name in reader.fieldnames if name}
    rows = list(reader)
    kalk = _parse_kalkulation_csv(rows, index)
    if kalk:
        return _with_it_costs(kalk)
    zeit = _parse_fb_it_csv(rows, index)
    if zeit:
        return _with_it_costs(zeit)
    phase_col = next((index[k] for k in index if k in {"phase", "abschnitt"}), None)
    party_col = next((index[k] for k in index if k in {"bereich", "partei", "team"}), None)
    pt_col = next((index[k] for k in index if k in {"pt", "personentage", "tage"}), None)
    cost_col = next(
        (index[k] for k in index if k in {"sachkosten", "kosten", "euro", "eur"}),
        None,
    )
    if not pt_col:
        raise EffortSheetError("Spalte PT fehlt.")
    totals = {
        "concept_scs": 0.0,
        "concept_cit": 0.0,
        "operate_scs": 0.0,
        "operate_cit": 0.0,
        "costs": 0.0,
    }
    for row in rows:
        if not any(str(value or "").strip() for value in row.values()):
            continue
        phase = _phase(row.get(phase_col) if phase_col else "")
        party = _party(row.get(party_col) if party_col else "")
        pt = parse_number(row.get(pt_col))
        totals[f"{phase}_{party}"] += pt
        if cost_col:
            totals["costs"] += parse_number(row.get(cost_col))
    concept_scs = totals["concept_scs"]
    concept_cit = totals["concept_cit"]
    operate_scs = totals["operate_scs"]
    operate_cit = totals["operate_cit"]
    costs = totals["costs"]
    return _with_it_costs({
        "effort_fb": _format_pt(concept_scs),
        "effort_it": _format_pt(concept_cit),
        "concept_scs_pt": _qty(concept_scs),
        "concept_cit_pt": _qty(concept_cit),
        "operate_scs_pt": _qty(operate_scs),
        "operate_cit_pt": _qty(operate_cit),
        "costs": _qty(costs) if costs else "",
    })


def _with_it_costs(parsed: dict[str, str]) -> dict[str, str]:
    rates = load_rates()
    cit_pt = parse_number(parsed.get("effort_it"))
    if not cit_pt:
        cit_pt = parse_number(parsed.get("concept_cit_pt")) + parse_number(
            parsed.get("operate_cit_pt")
        )
    cit_mat = parse_number(parsed.get("concept_cit_material")) + parse_number(
        parsed.get("operate_cit_material")
    )
    scs_pt = parse_number(parsed.get("concept_scs_pt")) + parse_number(
        parsed.get("operate_scs_pt")
    )
    if not scs_pt:
        scs_pt = parse_number(parsed.get("effort_fb"))
    scs_mat = parse_number(parsed.get("concept_scs_material")) + parse_number(
        parsed.get("operate_scs_material")
    )
    if not scs_mat:
        prev = parse_number(parsed.get("costs"))
        labor = scs_pt * rates.scs_daily
        if prev and (not labor or prev < labor):
            scs_mat = prev
    scs_total = round(_scs_total(scs_pt, scs_mat, rates), 2)
    cit_total = round(_cit_total(cit_pt, cit_mat, rates), 2)
    if scs_total:
        parsed["costs"] = _qty(scs_total)
    if not parsed.get("it_costs") and cit_total:
        parsed["it_costs"] = _qty(cit_total)
    return parsed


def _format_pt(value: float) -> str:
    if value <= 0:
        return "0 PT"
    if value == int(value):
        return f"{int(value)} PT"
    return f"{value:.1f}".replace(".", ",") + " PT"


def _qty(value: float) -> str:
    if value <= 0:
        return ""
    if value == int(value):
        return str(int(value))
    return str(value).replace(".", ",")


def fetch_effort_sheet(url: str, *, timeout: int = 15) -> dict[str, str]:
    if is_dummy_template_url(url):
        return parse_dummy_template(url)
    export = export_csv_url(url)
    try:
        response = httpx.get(
            export,
            timeout=timeout,
            follow_redirects=True,
            trust_env=False,
            headers={"accept": "text/csv,text/plain,*/*"},
        )
    except httpx.HTTPError as err:
        raise EffortSheetError("Sheet nicht erreichbar.") from err
    host = (urlparse(str(response.url)).hostname or "").lower()
    if host not in _ALLOWED_HOSTS and "googleusercontent.com" not in host:
        raise EffortSheetError(_NOT_SHARED)
    if response.status_code >= 400:
        raise EffortSheetError(_NOT_SHARED)
    ctype = (response.headers.get("content-type") or "").lower()
    body = response.text or ""
    if "html" in ctype or body.lstrip()[:15].lower().startswith(("<!doctype", "<html")):
        raise EffortSheetError(_NOT_SHARED)
    parsed = parse_effort_csv(body)
    parsed["effort_sheet_url"] = (url or "").strip()
    return parsed


EFFORT_SYNC_KEYS = (
    "effort_sheet_url",
    "costs",
    "it_costs",
    "concept_scs_pt",
    "concept_cit_pt",
    "operate_scs_pt",
    "operate_cit_pt",
    "concept_scs_pt_detail",
    "concept_cit_pt_detail",
    "operate_scs_pt_detail",
    "operate_cit_pt_detail",
    "concept_scs_material",
    "concept_cit_material",
    "operate_scs_material",
    "operate_cit_material",
)


def effort_sync_fields(parsed: dict[str, str]) -> dict[str, str]:
    return {key: str(parsed.get(key) or "").strip() for key in EFFORT_SYNC_KEYS}


def commit_effort_csv(
    db: Session,
    csv_text: str,
    public_base: str,
    share_id: str | None = None,
) -> dict[str, str]:
    parsed = parse_effort_csv(csv_text)
    sid = share_id_from_url(f"/aufwand/{share_id}") if share_id else None
    if not sid:
        sid = str(uuid.uuid4())
    key = f"{_SHARE_PREFIX}{sid}"
    row = db.get(AppSetting, key)
    if row:
        row.value = csv_text
        row.secret = False
    else:
        db.add(AppSetting(key=key, value=csv_text, secret=False))
    db.flush()
    parsed["effort_sheet_url"] = f"{public_base.rstrip('/')}/aufwand/{sid}"
    parsed["share_id"] = sid
    return parsed


def load_share_csv(db: Session, share_id: str) -> str | None:
    try:
        uuid.UUID(share_id)
    except ValueError:
        return None
    row = db.get(AppSetting, f"{_SHARE_PREFIX}{share_id}")
    return row.value if row and row.value else None


def kalkulation_xlsx(csv_text: str) -> bytes:
    from io import BytesIO

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    navy = "0A1A4F"
    cit_blue = "4A6BE8"
    head = "EEF3FF"
    line = "C9D6F4"
    thin = Border(
        left=Side(style="thin", color=line),
        right=Side(style="thin", color=line),
        top=Side(style="thin", color=line),
        bottom=Side(style="thin", color=line),
    )
    fills = {
        "phase": PatternFill("solid", fgColor=navy),
        "scs": PatternFill("solid", fgColor=navy),
        "cit": PatternFill("solid", fgColor=cit_blue),
        "colhead": PatternFill("solid", fgColor=head),
        "brutto": PatternFill("solid", fgColor=navy),
        "gesamt": PatternFill("solid", fgColor="F4F7FF"),
    }
    fonts = {
        "phase": Font(name="Calibri", size=14, bold=True, color="FFFFFF"),
        "scs": Font(name="Calibri", size=12, bold=True, color="FFFFFF"),
        "cit": Font(name="Calibri", size=12, bold=True, color="FFFFFF"),
        "colhead": Font(name="Calibri", size=10, bold=True, color=navy),
        "brutto": Font(name="Calibri", size=11, bold=True, color="FFFFFF"),
        "gesamt": Font(name="Calibri", size=11, bold=True, color=navy),
        "body": Font(name="Calibri", size=11, color=navy),
    }

    def kind(cells: list[str]) -> str:
        first = (cells[0] if cells else "").strip().lower()
        blob = " ".join(cells).lower()
        if first.startswith("konzeptions") or first.startswith("betriebs"):
            return "phase"
        if first.startswith("kalkulation it"):
            return "cit"
        if first.startswith("kalkulation scs"):
            return "scs"
        if first.startswith("personentage") or first.startswith("sach-"):
            return "colhead"
        if first.startswith("kosten-kalkulation"):
            return "colhead"
        if first.startswith("summe brutto"):
            return "brutto"
        if "gesamt-personentage" in blob:
            return "gesamt"
        return "body"

    book = Workbook()
    sheet = book.active
    sheet.title = "Kalkulation"
    sheet.sheet_view.showGridLines = False
    sheet.page_setup.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    widths = (22, 58, 16, 22)
    for idx, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(idx)].width = width

    rows = _csv_rows(csv_text)
    for r_idx, raw in enumerate(rows, start=1):
        cells = [(raw[i] if i < len(raw) else "") for i in range(4)]
        style = kind(cells)
        fill = fills.get(style)
        font = fonts.get(style, fonts["body"])
        for c_idx, value in enumerate(cells, start=1):
            cell = sheet.cell(r_idx, c_idx, value)
            cell.font = font
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            cell.border = thin
            if fill:
                cell.fill = fill
        if style in {"phase", "scs", "cit"}:
            sheet.merge_cells(start_row=r_idx, start_column=1, end_row=r_idx, end_column=4)
            sheet.row_dimensions[r_idx].height = 22
        elif style == "brutto":
            sheet.row_dimensions[r_idx].height = 20

    sheet.freeze_panes = "A2"
    buf = BytesIO()
    book.save(buf)
    return buf.getvalue()


def template_bytes() -> bytes:
    path = Path(TEMPLATE_PATH)
    return path.read_bytes()
